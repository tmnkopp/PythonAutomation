import re, sys, argparse, json, xlrd, openpyxl, os
import pandas as pd 
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from tabulate import tabulate
from context import context

class questionnaire_picklist_parser():
    def __init__(self, ctx):  
        self.ctx = ctx
        self.source = self.ctx.config['source']
    def parse(self):
        wb = openpyxl.load_workbook(filename = self.source)
        ws = wb.worksheets[2]
        lst=[]
        id=''
        indexes=self.ctx.config['indexes'] 
        for i,c in enumerate(ws['A']): 
            if c.value != None:
                if re.search(r'(CQ\d.*)',c.value,flags=re.I): 
                    id=c.value   
            if id!='' and ws[f'D{i+1}'].value is not None:
                d={}
                d['id']=id
                d['DisplayValue']=ws[f'D{i+1}'].value
                for di in indexes:
                    for key in di:
                        di[key]=di[key]+1
                        d[key]=di[key]
                lst.append(d) 
        df=pd.DataFrame(lst)  
        return df 
class questionnaire_parser():
    def __init__(self, ctx):  
        self.ctx = ctx
        self.source = self.ctx.config['source']
    def parse(self):
        wb = openpyxl.load_workbook(filename = self.source )
        ws = wb.worksheets[2]
        lst=[]
        irow=1
        indexes=self.ctx.config['indexes'] 
        for i,r in enumerate(ws['A']): 
            a=r.value 
            if a != None:
                if re.search(r'(CQ\d.*)', a,  flags=re.I): 
                    d={}
                    d['TEMPLNAME']=self._get_ftype(ws, irow)
                    d['HTMLID']=a
                    d['QTEXT']=ws[f'B{i+1}'].value  
                    d['CSSCLASS']=a 
                    for di in indexes:
                        for key in di:
                            di[key]=di[key]+1
                            d[key]=di[key]
                    lst.append(d)
                    irow=irow+1  
        df=pd.DataFrame(lst)  
        return df
    def _get_ftype(self, ws, irow):
        t='PICK'
        tval='' 
        if ws[f'D{irow}'].value is not None:
            tval=ws[f'D{irow}'].value 
        if re.search(r'Free text',tval,flags=re.I):
            t='TXT'
        return t     
 
class cogen_script():
    def __init__(self):  
        self.root = os.path.dirname(os.path.realpath(__file__)) 
        self.ext='.txt'
    def generate(self, ctx):
        df=ctx.payload
        st='' 
        for i,r in df.iterrows():
            tmp = ctx.get_template()
            if 'TEMPLNAME' in r.keys(): 
                tmp = tmp.replace('{TEMPLNAME}',r['TEMPLNAME'])  
            self.ext = tmp[tmp.index('.'):] 
            with open(tmp, 'r', encoding=ctx.config['encoding'], errors='replace') as f: 
                tmp = f.read()
            for c in df.columns:    
                tmp = tmp.replace("{"+c+"}", str(r[c]))   
            st=st+tmp   
        return st
 
def main(): 
    ctx=context() 
    parser = argparse.ArgumentParser(description="ArgumentParser"
        ,  formatter_class=argparse.ArgumentDefaultsHelpFormatter) 
    parser.add_argument("-v", "--verbose", action="store_true", help="increase verbosity") 
    args = vars(parser.parse_args())
    ctx.args=args 
    with open('cogen.config', 'r') as f: 
        ctx.config=json.loads(f.read()) 
     
    if ctx.args['verbose']:
        print(json.dumps(ctx.config, indent=1))
        print(json.dumps(ctx.args, indent=1))

    parser = globals()[ctx.config['parser']](ctx)   
    df=parser.parse()
    ctx.payload=df
    if ctx.config['parse_to']=='console':
        print(tabulate(df, headers = 'keys', tablefmt = 'plain')) 

    cogen = globals()['cogen_script']()   
    st=cogen.generate(ctx) 

    if ctx.config['generate_to']=='script':  
        with open(f'{ctx.get_dir()}\out{cogen.ext}', 'w', encoding=ctx.config['encoding']) as f: 
            f.write(st) 
    if ctx.config['generate_to']=='console':
        print(st)

if __name__ == "__main__": 
   main()
 
